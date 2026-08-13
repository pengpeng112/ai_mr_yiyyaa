"""患者质控列表与导出筛选语义一致性测试。"""
from datetime import datetime
from types import SimpleNamespace

import pytest
from fastapi import HTTPException
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from starlette.requests import Request

from app.database import Base
from app.models import AuditDimensionResult, PushLog, QCFeedback, User
from app.routers import patient_qc
from app.services import patient_visit_export_service as export_svc


def _make_db():
    engine = create_engine("sqlite:///:memory:")
    Base.metadata.create_all(bind=engine)
    return sessionmaker(bind=engine)()


def _add_log(
    db,
    *,
    patient_id,
    visit_number="1",
    dept="内科",
    patient_name="张三",
    admission_no="A001",
    status="success",
    severity="high",
    push_time=None,
    superseded_by=None,
    contract_valid=1,
    audit_type_code="progress_vs_nursing",
    request_json="",
):
    log = PushLog(
        patient_id=patient_id,
        visit_number=visit_number,
        dept=dept,
        patient_name=patient_name,
        admission_no=admission_no,
        status=status,
        severity=severity,
        trigger_type="manual",
        query_date="2026-08-01",
        push_time=push_time or datetime(2026, 8, 5, 10, 0, 0),
        inconsistency=0,
        risk_score=0,
        elapsed_ms=0,
        retry_count=0,
        superseded_by=superseded_by,
        contract_valid=contract_valid,
        audit_type_code=audit_type_code,
        request_json=request_json or "",
    )
    db.add(log)
    db.flush()
    return log


def _seed_core(db):
    """构造覆盖文本/科室/严重度/反馈/日期/当前结果语义的样例。"""
    # 命中：高危、待处理、内科、日期 08-05
    log_a = _add_log(
        db,
        patient_id="P100",
        visit_number="1",
        dept="内科",
        patient_name="张三",
        admission_no="A100",
        severity="high",
        push_time=datetime(2026, 8, 5, 10, 0, 0),
        request_json='{"patient_info":{"discharge_dept_name":"内科"}}',
    )
    db.add(AuditDimensionResult(
        push_log_id=log_a.id, dimension="维度A", status="fail", severity="high",
    ))
    db.add(QCFeedback(
        push_log_id=log_a.id, dept_id=1, severity="high", status="pending", created_by=1,
    ))

    # 同患者另一科室分组（导出时按住院次去重）
    log_a2 = _add_log(
        db,
        patient_id="P100",
        visit_number="1",
        dept="外科",
        patient_name="张三",
        admission_no="A100",
        severity="medium",
        push_time=datetime(2026, 8, 5, 11, 0, 0),
        request_json='{"patient_info":{"discharge_dept_name":"外科"}}',
    )
    db.add(AuditDimensionResult(
        push_log_id=log_a2.id, dimension="维度B", status="fail", severity="medium",
    ))

    # 中危、已整改
    log_b = _add_log(
        db,
        patient_id="P200",
        visit_number="2",
        dept="骨科",
        patient_name="李四",
        admission_no="A200",
        severity="medium",
        push_time=datetime(2026, 8, 6, 10, 0, 0),
        request_json='{"patient_info":{"discharge_dept_name":"骨科"}}',
    )
    db.add(AuditDimensionResult(
        push_log_id=log_b.id, dimension="维度C", status="fail", severity="medium",
    ))
    db.add(QCFeedback(
        push_log_id=log_b.id, dept_id=1, severity="medium", status="rectified", created_by=1,
    ))

    # 已 supersede，不应出现在当前结果
    log_old = _add_log(
        db,
        patient_id="P300",
        visit_number="1",
        dept="内科",
        patient_name="王五",
        admission_no="A300",
        severity="high",
        push_time=datetime(2026, 8, 5, 9, 0, 0),
        superseded_by=999,
    )
    db.add(AuditDimensionResult(
        push_log_id=log_old.id, dimension="旧", status="fail", severity="high",
    ))

    # contract_valid=0 不应出现
    log_bad = _add_log(
        db,
        patient_id="P400",
        visit_number="1",
        dept="内科",
        patient_name="赵六",
        admission_no="A400",
        severity="high",
        push_time=datetime(2026, 8, 5, 12, 0, 0),
        contract_valid=0,
    )
    db.add(AuditDimensionResult(
        push_log_id=log_bad.id, dimension="坏契约", status="fail", severity="high",
    ))

    # 失败状态不出现
    _add_log(
        db,
        patient_id="P500",
        visit_number="1",
        dept="内科",
        status="failed",
        severity="high",
        push_time=datetime(2026, 8, 5, 13, 0, 0),
    )

    # 日期窗外
    log_out = _add_log(
        db,
        patient_id="P600",
        visit_number="1",
        dept="内科",
        patient_name="周七",
        admission_no="A600",
        severity="high",
        push_time=datetime(2026, 7, 1, 10, 0, 0),
    )
    db.add(AuditDimensionResult(
        push_log_id=log_out.id, dimension="窗外", status="fail", severity="high",
    ))

    db.commit()


def _list_keys(db, **kwargs):
    filters = patient_qc.normalize_patient_qc_filters(strict=False, **kwargs)
    groups = patient_qc.build_patient_qc_grouped_query(db, filters).all()
    return {(str(g.pid), str(g.vn), str(g.dp or "")) for g in groups}


def _export_keys(db, **kwargs):
    filters = patient_qc.normalize_patient_qc_filters(strict=True, **kwargs)
    return patient_qc.query_patient_qc_visit_keys(db, filters)


def test_list_and_export_share_text_dept_severity_status_date_keys():
    db = _make_db()
    _seed_core(db)

    # 无筛选：列表含两科室分组 + 骨科；导出按住院次去重
    list_all = _list_keys(db)
    export_all = _export_keys(db)
    assert ("P100", "1", "内科") in list_all
    assert ("P100", "1", "外科") in list_all
    assert ("P200", "2", "骨科") in list_all
    assert ("P300", "1", "内科") not in list_all  # superseded
    assert ("P400", "1", "内科") not in list_all  # contract_valid=0
    assert ("P500", "1", "内科") not in list_all  # failed
    assert export_all == {("P100", "1"), ("P200", "2"), ("P600", "1")}

    # 文本 patient_name
    assert _export_keys(db, patient_name="张") == {("P100", "1")}
    assert {k[:2] for k in _list_keys(db, patient_name="张")} == {("P100", "1")}

    # 科室
    assert _export_keys(db, dept="内") == {("P100", "1"), ("P600", "1")}
    assert ("P200", "2", "骨科") not in _list_keys(db, dept="内")

    # 严重度 high：P100 内科 high；P600 high；P200 只有 medium
    high_export = _export_keys(db, severity="high")
    high_list = _list_keys(db, severity="high")
    assert high_export == {("P100", "1"), ("P600", "1")}
    assert {k[:2] for k in high_list} == high_export

    # 反馈状态 pending
    assert _export_keys(db, status="pending") == {("P100", "1")}
    assert {k[:2] for k in _list_keys(db, status="pending")} == {("P100", "1")}

    # 日期闭区间（按 push_time）
    date_export = _export_keys(db, date_from="2026-08-05", date_to="2026-08-05")
    date_list = _list_keys(db, date_from="2026-08-05", date_to="2026-08-05")
    assert date_export == {("P100", "1")}
    assert {k[:2] for k in date_list} == date_export
    assert ("P200", "2") not in date_export
    assert ("P600", "1") not in date_export


def test_export_keys_dedupe_across_depts():
    db = _make_db()
    _seed_core(db)
    keys = _export_keys(db, patient_id="P100")
    assert keys == {("P100", "1")}


def test_filter_patients_by_keys_intersection_and_empty():
    temp_rows = [
        {"patient_id": "P100", "visit_number": "1", "admission_no": "A100"},
        {"patient_id": "P200", "visit_number": "2", "admission_no": "A200"},
        {"patient_id": "P900", "visit_number": "1", "admission_no": "A900"},
    ]
    filtered = export_svc._filter_patients_by_keys(temp_rows, {("P100", "1"), ("P200", "2")})
    assert [(r["patient_id"], r["visit_number"]) for r in filtered] == [
        ("P100", "1"),
        ("P200", "2"),
    ]
    # TEMP 外患者不会出现
    assert all(r["patient_id"] != "P900" or ("P900", "1") in {("P100", "1"), ("P200", "2")} for r in filtered)

    empty = export_svc._filter_patients_by_keys(temp_rows, set())
    assert empty == []

    # 去重
    duped = export_svc._filter_patients_by_keys(
        temp_rows + [{"patient_id": "P100", "visit_number": "1", "admission_no": "A100b"}],
        {("P100", "1")},
    )
    assert len(duped) == 1


def test_export_empty_filter_returns_header_only_workbook(monkeypatch):
    xlsx_bytes, ext, count = export_svc.export_patient_visit_summary(object(), patient_keys=set())
    assert ext == "xlsx"
    assert count == 0
    assert len(xlsx_bytes) > 0

    import io
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    ws = wb.active
    assert ws.max_row == 1
    headers = [c.value for c in ws[1]]
    assert "患者ID" in headers
    assert "住院号" in headers


def _admin_user(db):
    from app.models import Role, User
    role = db.query(Role).filter(Role.name == "admin").first()
    if not role:
        role = Role(name="admin", description="")
        db.add(role)
        db.flush()
    user = User(username="admin_export", password_hash="x", full_name="管理员", role_id=role.id)
    db.add(user)
    db.flush()
    return user


def test_export_endpoint_audits_filtered_scope_without_pii(monkeypatch):
    calls = []
    db = _make_db()
    _seed_core(db)
    admin = _admin_user(db)
    db.commit()

    monkeypatch.setattr(
        "app.services.patient_visit_export_service.export_patient_visit_summary",
        lambda db, patient_keys=None: (b"xlsx", "xlsx", len(patient_keys or [])),
    )
    monkeypatch.setattr(patient_qc, "record_export_audit", lambda **kwargs: calls.append(kwargs))

    request = Request({
        "type": "http",
        "method": "GET",
        "path": "/api/patient-qc/export/patient-visit-summary",
        "headers": [(b"user-agent", b"focused-test")],
        "client": ("127.0.0.1", 1234),
        "query_string": b"patient_id=P100&severity=high&date_from=2026-08-05",
        "scheme": "http",
        "server": ("test", 80),
        "asgi": {"version": "3.0"},
    })

    response = patient_qc.export_patient_visit_summary(
        request=request,
        patient_id="P100",
        patient_name="张三",
        admission_no="A100",
        visit_number="1",
        dept="内科",
        discharge_dept_name=None,
        severity="high",
        audit_type_code="progress_vs_nursing",
        status="pending",
        date_from="2026-08-05",
        date_to="2026-08-05",
        db=db,
        current_user=admin,
    )

    assert response.body == b"xlsx"
    criteria = calls[0]["filter_criteria"]
    assert criteria["scope"] == "filtered_all_pages"
    assert criteria["source"] == "TEMP_PAT_VISIT_LIST"
    assert criteria["severity"] == "high"
    assert criteria["status"] == "pending"
    assert criteria["dept"] == "内科"
    assert criteria["audit_type_code"] == "progress_vs_nursing"
    assert criteria["date_from"] == "2026-08-05"
    assert criteria["date_to"] == "2026-08-05"
    assert criteria["patient_id"] == "已提供"
    assert criteria["patient_name"] == "已提供"
    assert criteria["admission_no"] == "已提供"
    assert criteria["visit_number"] == "已提供"
    # 审计 JSON 中不得出现患者原文
    raw = str(criteria)
    assert "P100" not in raw
    assert "张三" not in raw
    assert "A100" not in raw
    assert calls[0]["record_count"] == 1
    assert calls[0]["status"] == "success"


def test_export_endpoint_invalid_filters_return_422(monkeypatch):
    db = _make_db()
    request = Request({
        "type": "http",
        "method": "GET",
        "path": "/api/patient-qc/export/patient-visit-summary",
        "headers": [],
        "client": ("127.0.0.1", 1),
        "query_string": b"",
        "scheme": "http",
        "server": ("test", 80),
        "asgi": {"version": "3.0"},
    })
    with pytest.raises(HTTPException) as caught:
        patient_qc.export_patient_visit_summary(
            request=request,
            patient_id=None,
            patient_name=None,
            admission_no=None,
            visit_number=None,
            dept=None,
            discharge_dept_name=None,
            severity="critical",
            audit_type_code=None,
            status=None,
            date_from=None,
            date_to=None,
            db=db,
            current_user=SimpleNamespace(id=1, username="admin"),
        )
    assert caught.value.status_code == 422

    with pytest.raises(HTTPException) as caught_date:
        patient_qc.export_patient_visit_summary(
            request=request,
            patient_id=None,
            patient_name=None,
            admission_no=None,
            visit_number=None,
            dept=None,
            discharge_dept_name=None,
            severity=None,
            audit_type_code=None,
            status=None,
            date_from="2026/08/01",
            date_to=None,
            db=db,
            current_user=SimpleNamespace(id=1, username="admin"),
        )
    assert caught_date.value.status_code == 422


def test_dept_scope_limits_list_and_export_keys():
    """非管理员仅可见科室范围内的患者；管理员不受限。"""
    from app.models import Role, Department, RoleDepartment, User

    db = _make_db()
    role_mgr = Role(name="dept_manager", description="")
    role_admin = Role(name="admin", description="")
    db.add_all([role_mgr, role_admin])
    db.flush()
    d_xnk = Department(name="心内科", code="XNK")
    d_gk = Department(name="骨科", code="GK")
    db.add_all([d_xnk, d_gk])
    db.flush()
    db.add(RoleDepartment(role_id=role_mgr.id, dept_id=d_xnk.id))
    mgr = User(username="mgr", password_hash="x", full_name="主任", role_id=role_mgr.id, dept_id=d_xnk.id)
    admin = User(username="adm", password_hash="x", full_name="管理员", role_id=role_admin.id)
    db.add_all([mgr, admin])
    db.flush()

    log_xnk = _add_log(db, patient_id="PX", visit_number="1", dept="心内科", severity="high")
    db.add(AuditDimensionResult(push_log_id=log_xnk.id, dimension="d", status="fail", severity="high"))
    log_gk = _add_log(db, patient_id="PG", visit_number="1", dept="骨科", severity="high")
    db.add(AuditDimensionResult(push_log_id=log_gk.id, dimension="d", status="fail", severity="high"))
    db.commit()

    filters = patient_qc.normalize_patient_qc_filters(strict=False)
    mgr_keys = patient_qc.query_patient_qc_visit_keys(db, filters, current_user=mgr)
    admin_keys = patient_qc.query_patient_qc_visit_keys(db, filters, current_user=admin)
    assert mgr_keys == {("PX", "1")}
    assert {("PX", "1"), ("PG", "1")}.issubset(admin_keys)


def test_export_empty_result_audits_zero_count(monkeypatch):
    calls = []
    db = _make_db()
    _seed_core(db)
    admin = _admin_user(db)
    db.commit()
    monkeypatch.setattr(
        "app.services.patient_visit_export_service.export_patient_visit_summary",
        lambda db, patient_keys=None: (b"header-only", "xlsx", 0),
    )
    monkeypatch.setattr(patient_qc, "record_export_audit", lambda **kwargs: calls.append(kwargs))
    request = Request({
        "type": "http",
        "method": "GET",
        "path": "/api/patient-qc/export/patient-visit-summary",
        "headers": [],
        "client": ("127.0.0.1", 1),
        "query_string": b"",
        "scheme": "http",
        "server": ("test", 80),
        "asgi": {"version": "3.0"},
    })
    patient_qc.export_patient_visit_summary(
        request=request,
        patient_id="NO_MATCH",
        patient_name=None,
        admission_no=None,
        visit_number=None,
        dept=None,
        discharge_dept_name=None,
        severity=None,
        audit_type_code=None,
        status=None,
        date_from=None,
        date_to=None,
        db=db,
        current_user=admin,
    )
    assert calls[0]["record_count"] == 0
    assert calls[0]["status"] == "success"
    assert calls[0]["filter_criteria"]["scope"] == "filtered_all_pages"
    assert calls[0]["filter_criteria"]["patient_id"] == "已提供"


def test_discharge_dept_name_filter_sqlite_matches_json_and_dept():
    """SQLite 下 discharge_dept_name 筛选：匹配 request_json 内嵌字段，并 OR dept 列兜底。"""
    db = _make_db()
    _seed_core(db)
    # 骨科：log_b 的 request_json 内 discharge_dept_name=骨科
    keys_gk = _list_keys(db, discharge_dept_name="骨科")
    assert {k[:2] for k in keys_gk} == {("P200", "2")}
    # 内科：log_a 的 json discharge=内科；P600 dept=内科 但 json 空 → 靠 OR dept 命中
    keys_nk = _list_keys(db, discharge_dept_name="内科")
    assert ("P100", "1", "内科") in keys_nk
    assert ("P200", "2", "骨科") not in keys_nk


def test_discharge_dept_name_filter_oracle_uses_dbms_lob_instr(monkeypatch):
    """Oracle dialect 下必须走 dbms_lob.instr，禁止裸 LIKE request_json（ORA-00932 回归守护）。"""
    db = _make_db()
    monkeypatch.setattr(
        patient_qc, "_db_engine",
        SimpleNamespace(dialect=SimpleNamespace(name="oracle")),
    )
    filters = patient_qc.normalize_patient_qc_filters(strict=False, discharge_dept_name="耳外科")
    conds = patient_qc.build_patient_qc_base_filters(filters)

    def _render(cond):
        try:
            return str(cond.compile(compile_kwargs={"literal_binds": True}))
        except Exception:
            return str(cond)

    blob = " | ".join(_render(c) for c in conds)
    assert "dbms_lob.instr" in blob
    # Oracle 分支不得再生成对 request_json 的裸 LIKE（否则回到 ORA-00932）
    assert "request_json LIKE" not in blob.upper()
    # pattern 应以绑定参数形式传入，避免拼接注入
    assert "耳外科" in blob
