"""患者质控结果汇总导出（应用库数据源，支持历史出院患者）。

与患者质控列表同源同筛选；每位患者一行，附高危/中危问题明细与整改建议。
不依赖 Oracle TEMP_PAT_VISIT_LIST，筛选命中即可导出。
"""
from __future__ import annotations

import io
from datetime import datetime
from typing import Optional

from sqlalchemy import or_
from sqlalchemy.orm import Session

from app.models import AuditDimensionResult, PushLog, QCFeedback

_EXCEL_CELL_MAX = 32767
_CHUNK_GROUPS = 200
_CHUNK_IDS = 900

# (列名, 列宽)
QC_SUMMARY_COLUMNS = [
    ("患者ID", 14), ("患者姓名", 10), ("住院次", 8), ("住院号", 12),
    ("在院科室", 14), ("出院科室", 14), ("最高严重度", 10), ("预警级别", 10),
    ("高危数", 8), ("中危数", 8), ("低危数", 8), ("问题数", 8),
    ("待处理数", 9), ("已闭环数", 9), ("审计类型数", 10), ("推送次数", 9),
    ("最近推送时间", 17), ("问题明细", 60), ("整改建议", 50),
]

_SEVERITY_RANK = {"high": 3, "medium": 2, "low": 1, "": 0, None: 0}


def _is_issue(dim) -> bool:
    status = str(getattr(dim, "status", "") or "").lower()
    if status in {"fail", "warn", "warning", "risk"}:
        return True
    return bool(str(getattr(dim, "issue_summary", "") or "").strip())


def _alert_level_for_severity(severity: str) -> str:
    return {"high": "red", "medium": "yellow", "low": "blue"}.get(severity, "")


def _format_dt(dt) -> str:
    if not dt:
        return ""
    if isinstance(dt, datetime):
        return dt.strftime("%Y-%m-%d %H:%M:%S")
    return str(dt)[:19]


def _snapshot_name(log) -> tuple[str, str]:
    """从 request_json 解析患者姓名/住院号（失败回退 PushLog 字段）。"""
    import json
    name = ""
    admission_no = ""
    raw = getattr(log, "request_json", None)
    if raw:
        try:
            data = json.loads(raw) if isinstance(raw, str) else raw
            info = data.get("patient_info", {}) if isinstance(data, dict) else {}
            if isinstance(info, dict):
                name = str(info.get("patient_name") or "")
                admission_no = str(info.get("admission_no") or "")
        except (TypeError, ValueError):
            pass
    return (name or getattr(log, "patient_name", "") or "",
            admission_no or getattr(log, "admission_no", "") or "")


def build_qc_summary_rows(groups, db: Session) -> list[list[str]]:
    """把分组查询结果组装为 Excel 行。"""
    logs_by_group: dict[str, list] = {}
    group_keys = [(g.pid, g.vn, g.dp) for g in groups]
    all_log_ids: list[int] = []
    for start in range(0, len(group_keys), _CHUNK_GROUPS):
        batch = group_keys[start:start + _CHUNK_GROUPS]
        conds = [
            (PushLog.patient_id == pid) & (PushLog.visit_number == vn) & (PushLog.dept == dp)
            for pid, vn, dp in batch
        ]
        for log in db.query(PushLog).filter(PushLog.status == "success", or_(*conds)).all():
            key = f"{log.patient_id}::{log.visit_number}::{log.dept}"
            logs_by_group.setdefault(key, []).append(log)
            all_log_ids.append(log.id)

    dims_by_log: dict[int, list] = {}
    fb_by_log: dict[int, list] = {}
    for start in range(0, len(all_log_ids), _CHUNK_IDS):
        ids = all_log_ids[start:start + _CHUNK_IDS]
        for d in db.query(AuditDimensionResult).filter(AuditDimensionResult.push_log_id.in_(ids)).all():
            dims_by_log.setdefault(d.push_log_id, []).append(d)
        for f in db.query(QCFeedback).filter(QCFeedback.push_log_id.in_(ids)).all():
            fb_by_log.setdefault(f.push_log_id, []).append(f)

    rows: list[list[str]] = []
    for g in groups:
        key = f"{g.pid}::{g.vn}::{g.dp}"
        g_logs = logs_by_group.get(key, [])
        if not g_logs:
            continue
        g_dims = [d for lid in [l.id for l in g_logs] for d in dims_by_log.get(lid, [])]
        g_fbs = [f for lid in [l.id for l in g_logs] for f in fb_by_log.get(lid, [])]

        high = [d for d in g_dims if getattr(d, "severity", "") == "high"]
        medium = [d for d in g_dims if getattr(d, "severity", "") == "medium"]
        low_n = sum(1 for d in g_dims if getattr(d, "severity", "") == "low")
        issue_n = sum(1 for d in g_dims if _is_issue(d))
        pending_n = sum(1 for f in g_fbs if getattr(f, "status", "") == "pending")
        resolved_n = sum(1 for f in g_fbs if getattr(f, "status", "") in {"rectified", "closed"})

        severities = [d.severity for d in g_dims if getattr(d, "severity", "")]
        highest = max(severities, key=lambda s: _SEVERITY_RANK.get(s, 0)) if severities else ""

        first_log = g_logs[0]
        name, admission_no = _snapshot_name(first_log)

        focus = sorted(high + medium, key=lambda d: -_SEVERITY_RANK.get(getattr(d, "severity", ""), 0))
        issue_lines: list[str] = []
        advice_lines: list[str] = []
        for d in focus:
            summary_text = str(getattr(d, "issue_summary", "") or getattr(d, "explanation", "") or "").strip()
            if not summary_text:
                continue
            dim_name = str(getattr(d, "dimension", "") or getattr(d, "dimension_code", "") or "").strip()
            prefix = f"[{d.severity}]{'[' + dim_name + ']' if dim_name else ''}"
            issue_lines.append(f"{prefix}{summary_text}")
            rec = str(getattr(d, "recommendation", "") or "").strip()
            if rec:
                advice_lines.append(f"{prefix}{rec}")

        rows.append([
            str(g.pid or ""), name, str(g.vn or ""), admission_no,
            str(g.dp or first_log.dept or ""), "",
            highest, _alert_level_for_severity(highest),
            len(high), len(medium), low_n, issue_n,
            pending_n, resolved_n,
            int(g.audit_type_count or 0), int(g.push_log_count or 0),
            _format_dt(g.latest_push_time),
            "\n".join(issue_lines)[:_EXCEL_CELL_MAX],
            "\n".join(advice_lines)[:_EXCEL_CELL_MAX],
        ])
    return rows


def build_qc_summary_excel(rows: list[list[str]]) -> bytes:
    """生成患者质控结果汇总 Excel。"""
    try:
        import openpyxl
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    except ImportError as exc:
        raise RuntimeError("openpyxl 未安装") from exc

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "患者质控结果汇总"
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=10)
    thin = Border(left=Side(style="thin"), right=Side(style="thin"),
                  top=Side(style="thin"), bottom=Side(style="thin"))
    for col_idx, (name, width) in enumerate(QC_SUMMARY_COLUMNS, 1):
        cell = ws.cell(row=1, column=col_idx, value=name)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin
        ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = width
    data_font = Font(size=9)
    data_align = Alignment(horizontal="left", vertical="top", wrap_text=True)
    for row_idx, row_data in enumerate(rows, 2):
        for col_idx, value in enumerate(row_data, 1):
            text = "" if value is None else str(value)
            cell = ws.cell(row=row_idx, column=col_idx, value=text[:_EXCEL_CELL_MAX])
            cell.font = data_font
            cell.alignment = data_align
            cell.border = thin
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
